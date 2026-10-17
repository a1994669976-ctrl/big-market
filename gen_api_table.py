# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "API清单"

# 表头
headers = ["层级", "接口/类名", "方法签名", "作用说明", "所在文件"]
ws.append(headers)

# 数据
rows = [
    # 一、领域服务 API
    ["领域服务", "IStartegyArmory", "void assembleLotteryStrategy(Long strategyId)",
     "策略装配：从DB读奖品概率→构建概率查找表→存入Redis",
     "domain/strategy/service/armory/IStartegyArmory.java"],
    ["领域服务", "IStartegyArmory", "Integer getRandomAwareId(Long strategyId)",
     "随机抽奖：在Redis概率表中按随机数随机出一个奖品ID",
     "domain/strategy/service/armory/IStartegyArmory.java"],

    # 二、仓储 API
    ["仓储接口", "IStrategyRepository", "List<StrategyAwardEntity> queryStrategyAwardList(Long strategyId)",
     "按策略ID查奖品列表（先查Redis缓存，miss再查MySQL，查到后回填Redis）",
     "domain/strategy/repository/IStrategyRepository.java"],
    ["仓储接口", "IStrategyRepository", "void storeStrategyAwardSearchRateTables(Long strategyId, BigDecimal rateRange, HashMap<Integer,Integer> tables)",
     "把构建好的概率查找表存入Redis（存两样：总范围值 + Map查找表）",
     "domain/strategy/repository/IStrategyRepository.java"],
    ["仓储接口", "IStrategyRepository", "int gatRateRange(Long strategyId)",
     "从Redis取概率总范围值（如10000，用于生成随机数上限）",
     "domain/strategy/repository/IStrategyRepository.java"],
    ["仓储接口", "IStrategyRepository", "Integer getStrategyAwardAssemble(Long strategyId, int rateKey)",
     "从Redis的Map里按随机数(rateKey)查对应的奖品ID",
     "domain/strategy/repository/IStrategyRepository.java"],

    # 三、DAO 层
    ["DAO(MySQL)", "IStrategyDao", "List<Strategy> queryStrategyList()",
     "查所有抽奖策略列表",
     "infrastructure/persistent/dao/IStrategyDao.java"],
    ["DAO(MySQL)", "IStrategyAwardDao", "List<StrategyAward> queryStrategyAwardList()",
     "查所有策略奖品配置",
     "infrastructure/persistent/dao/IStrategyAwardDao.java"],
    ["DAO(MySQL)", "IStrategyAwardDao", "List<StrategyAward> queryStrategyAwardListByStrategyId(Long strategyId)",
     "按策略ID查该策略下所有奖品配置",
     "infrastructure/persistent/dao/IStrategyAwardDao.java"],
    ["DAO(MySQL)", "IStrategyRuleDao", "List<StrategyRule> queryStrategyRuleList()",
     "查所有策略规则配置",
     "infrastructure/persistent/dao/IStrategyRuleDao.java"],
    ["DAO(MySQL)", "IAwardDao", "List<Award> queryAwardList()",
     "查所有奖品基础信息",
     "infrastructure/persistent/dao/IAwardDao.java"],

    # 四、Redis 服务
    ["Redis服务", "IRedisService", "<T> void setValue(String key, T value)",
     "存值到Redis",
     "infrastructure/persistent/redis/IRedisService.java"],
    ["Redis服务", "IRedisService", "<T> T getValue(String key)",
     "从Redis取值",
     "infrastructure/persistent/redis/IRedisService.java"],
    ["Redis服务", "IRedisService", "<K,V> RMap<K,V> getMap(String key)",
     "获取Redis Map结构",
     "infrastructure/persistent/redis/IRedisService.java"],
    ["Redis服务", "IRedisService", "String getFromMap(String key, String field)",
     "从Map中取指定字段的值",
     "infrastructure/persistent/redis/IRedisService.java"],
    ["Redis服务", "IRedisService", "Boolean setNx(String key, long expired, TimeUnit unit)",
     "分布式锁（SETNX）",
     "infrastructure/persistent/redis/IRedisService.java"],
    ["Redis服务", "IRedisService", "RLock getLock(String key)",
     "获取可重入锁",
     "infrastructure/persistent/redis/IRedisService.java"],
    ["Redis服务", "IRedisService", "RLock getFairLock(String key)",
     "获取公平锁",
     "infrastructure/persistent/redis/IRedisService.java"],
    ["Redis服务", "IRedisService", "RSemaphore getSemaphore(String key)",
     "获取信号量",
     "infrastructure/persistent/redis/IRedisService.java"],
    ["Redis服务", "IRedisService", "<T> RBloomFilter<T> getBloomFilter(String key)",
     "获取布隆过滤器",
     "infrastructure/persistent/redis/IRedisService.java"],

    # 五、公共组件
    ["公共组件", "Constants.RedisKey", "STRATEGY_AWARD_KEY = \"big_market_strategy_award_key_\"",
     "策略奖品列表缓存key前缀",
     "types/common/Constants.java"],
    ["公共组件", "Constants.RedisKey", "STRATEGY_RATE_TABLE_KEY = \"big_market_strategy_rate_table_key_\"",
     "概率查找表Map的key前缀",
     "types/common/Constants.java"],
    ["公共组件", "Constants.RedisKey", "STRATEGY_RATE_RANGE_KEY = \"big_market_strategy_rate_range_key_\"",
     "概率总范围值的key前缀",
     "types/common/Constants.java"],
    ["公共组件", "ResponseCode", "SUCCESS(\"0000\",\"成功\")",
     "响应码：成功",
     "types/enums/ResponseCode.java"],
    ["公共组件", "ResponseCode", "UN_ERROR(\"0001\",\"未知失败\")",
     "响应码：未知失败",
     "types/enums/ResponseCode.java"],
    ["公共组件", "ResponseCode", "ILLEGAL_PARAMETER(\"0002\",\"非法参数\")",
     "响应码：非法参数",
     "types/enums/ResponseCode.java"],
    ["公共组件", "AppException", "AppException(String code, String info)",
     "业务异常类（code+info）",
     "types/exception/AppException.java"],
]

for r in rows:
    ws.append(r)

# ===== 样式 =====
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

# 表头样式
for col in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = thin_border

# 数据行样式 + 层级分组着色
layer_colors = {
    "领域服务": "D9E2F3",
    "仓储接口": "E2EFDA",
    "DAO(MySQL)": "FFF2CC",
    "Redis服务": "FCE4D6",
    "公共组件": "EDEDED",
}
for row_idx in range(2, ws.max_row+1):
    layer = ws.cell(row=row_idx, column=1).value
    fill_color = layer_colors.get(layer, "FFFFFF")
    for col in range(1, len(headers)+1):
        cell = ws.cell(row=row_idx, column=col)
        cell.border = thin_border
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

# 列宽
ws.column_dimensions['A'].width = 12
ws.column_dimensions['B'].width = 22
ws.column_dimensions['C'].width = 55
ws.column_dimensions['D'].width = 50
ws.column_dimensions['E'].width = 50

# 冻结首行
ws.freeze_panes = "A2"

# 保存
output = r"E:\download\文档\学习\找实习\大营销\1\大营销项目API清单.xlsx"
wb.save(output)
print(f"已保存: {output}")
print(f"共 {ws.max_row-1} 条记录")
